"""The common tool suite, fully offline. Tools are exercised directly
(they are plain callables under ``@strands.tool``) and once end-to-end
through the scripted agent loop to prove name-based resolution."""

from __future__ import annotations

import json
import math
from datetime import date
from types import SimpleNamespace
from zoneinfo import ZoneInfoNotFoundError

import docx
import pytest
import requests
from strands.agent.state import AgentState

from talkpipe_agents import (
    AgentChat,
    LLMAgent,
    agent_tool_names,
    calculator,
    common_tools,
    current_time,
    days_between,
    fetch_url,
    list_files,
    read_file,
    regex_extract,
    search_files,
)


def test_suite_is_registered_by_name():
    assert {
        "calculator",
        "current_time",
        "days_between",
        "regex_extract",
        "read_file",
        "list_files",
        "search_files",
        "fetch_url",
        "save_note",
        "read_notes",
        "summarize_notes",
        "update_note",
        "delete_note",
        "clear_notes",
    } <= set(agent_tool_names())


# --- calculator -------------------------------------------------------


@pytest.mark.parametrize(
    ("expression", "expected"),
    [
        ("(3 + 4) * 2", "14"),
        ("2 ** 10", "1024"),
        ("7 // 2 + 7 % 2", "4"),
        ("-3 + +5", "2"),
        ("sqrt(16) + log10(100)", "6"),
        ("sin(pi / 6)", str(math.sin(math.pi / 6))),
        ("round(2.675, 3)", "2.675"),
        ("min(3, max(1, 2))", "2"),
        ("factorial(5)", "120"),
    ],
)
def test_calculator(expression, expected):
    assert calculator(expression) == expected


def test_calculator_whole_floats_shown_as_integers():
    assert calculator("10 / 2") == "5"


@pytest.mark.parametrize(
    "expression",
    [
        "__import__('os').system('true')",
        "().__class__",
        "x + 1",
        "lambda: 1",
        "1; 2",
        "9 ** 9 ** 9",
        "factorial(100000)",
        "'a' * 3",
    ],
)
def test_calculator_rejects_non_arithmetic(expression):
    with pytest.raises(ValueError):
        calculator(expression)


# --- dates and times --------------------------------------------------


def test_current_time_is_iso_with_weekday():
    stamp = current_time("UTC")
    assert "+00:00" in stamp
    day = date.fromisoformat(stamp[:10])
    assert f"({day:%A})" in stamp


def test_current_time_rejects_bad_timezone():
    with pytest.raises(ZoneInfoNotFoundError):
        current_time("Nowhere/Special")


def test_days_between():
    assert days_between("2026-01-31", "2026-03-01") == "29"
    assert days_between("2026-03-01", "2026-01-31") == "-29"
    with pytest.raises(ValueError, match="ISO"):
        days_between("31/01/2026", "2026-03-01")


# --- text -------------------------------------------------------------


def test_regex_extract_whole_matches_and_groups():
    assert (
        regex_extract(r"\b[A-Z]{2,5}\b", "Buy AMZN and MSFT, not tulips")
        == "AMZN\nMSFT"
    )
    assert regex_extract(r"(\w+)=(\d+)", "a=1 b=2") == "a\t1\nb\t2"
    assert regex_extract(r"\d+", "no digits here") == "(no matches)"
    assert regex_extract(r"\d", "123456", max_matches=2) == "1\n2"
    with pytest.raises(ValueError, match="invalid regular expression"):
        regex_extract("(", "text")


# --- files ------------------------------------------------------------


def test_read_file_truncates_with_notice(tmp_path):
    f = tmp_path / "long.txt"
    f.write_text("x" * 50)
    assert read_file(str(f)) == "x" * 50
    clipped = read_file(str(f), max_chars=10)
    assert clipped.startswith("x" * 10)
    assert "40 more characters" in clipped


def test_list_files(tmp_path):
    (tmp_path / "sub").mkdir()
    (tmp_path / "a.py").write_text("pass")
    (tmp_path / "b.txt").write_text("hi")
    listing = list_files(str(tmp_path))
    assert "sub/" in listing
    assert "a.py\t4 bytes" in listing
    assert list_files(str(tmp_path), "*.py").splitlines() == ["a.py\t4 bytes"]
    with pytest.raises(ValueError, match="not a directory"):
        list_files(str(tmp_path / "nope"))


def test_search_files(tmp_path):
    (tmp_path / "pkg").mkdir()
    (tmp_path / "pkg" / "mod.py").write_text("x = 1\nneedle here\n")
    (tmp_path / ".hidden").mkdir()
    (tmp_path / ".hidden" / "secret.py").write_text("needle hidden\n")
    (tmp_path / "blob.bin").write_bytes(b"needle\xff\xfe\x00")
    hits = search_files("needle", str(tmp_path))
    assert hits == "pkg/mod.py:2: needle here"
    assert search_files("absent", str(tmp_path)) == "(no matches)"


def test_search_files_respects_match_limit(tmp_path):
    (tmp_path / "many.txt").write_text("hit\n" * 10)
    hits = search_files("hit", str(tmp_path), max_matches=3)
    assert hits.count("many.txt") == 3
    assert "match limit reached" in hits


def test_read_file_understands_documents(tmp_path):
    """read_file is talkpipe's readFile: formats beyond plain text are
    read by their extractors, and anything else is read as text."""
    (tmp_path / "page.html").write_text(
        "<html><head><script>var x=1;</script></head>"
        "<body><h1>Title</h1><p>Hello <b>world</b></p></body></html>"
    )
    page = read_file(str(tmp_path / "page.html"))
    assert "Title" in page
    assert "Hello world" in page
    assert "var x" not in page
    assert "<p>" not in page
    (tmp_path / "rows.csv").write_text("name,n\nalpha,1\nbeta,2\n")
    assert read_file(str(tmp_path / "rows.csv")).splitlines() == [
        "name: alpha, n: 1",
        "name: beta, n: 2",
    ]
    doc = docx.Document()
    doc.add_paragraph("A Word paragraph.")
    doc.save(str(tmp_path / "memo.docx"))
    assert read_file(str(tmp_path / "memo.docx")) == "A Word paragraph."
    (tmp_path / "script.py").write_text("print('code')\n")
    assert read_file(str(tmp_path / "script.py")) == "print('code')\n"
    (tmp_path / "noext").write_text("bare")
    assert read_file(str(tmp_path / "noext")) == "bare"


def test_read_file_errors_reach_the_agent(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_file(str(tmp_path / "missing.txt"))
    (tmp_path / "blob.bin").write_bytes(b"\xff\xfe\x00")
    with pytest.raises(UnicodeDecodeError):
        read_file(str(tmp_path / "blob.bin"))


def test_search_files_searches_documents_by_their_text(tmp_path):
    (tmp_path / "page.html").write_text("<p>no</p><p>needle <b>in</b> html</p>")
    doc = docx.Document()
    doc.add_paragraph("a needle in Word")
    doc.save(str(tmp_path / "memo.docx"))
    (tmp_path / "rows.csv").write_text("name\nhay\nneedle\n")
    assert search_files("needle", str(tmp_path)).splitlines() == [
        "memo.docx:1: a needle in Word",
        "page.html:2: needle in html",
        "rows.csv:2: name: needle",
    ]


# --- web --------------------------------------------------------------
#
# fetch_url is talkpipe's downloadURL, which fetches with requests (and
# consults robots.txt first), so the network is faked at requests.get.


def _fake_requests(monkeypatch, body: bytes, content_type: str) -> list[str]:
    seen: list[str] = []

    def get(url, headers=None, timeout=None, **_):
        seen.append(url)
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = content_type
        response._content = body
        return response

    monkeypatch.setattr("talkpipe.data.html.requests.get", get)
    return seen


def test_fetch_url_extracts_html_text(monkeypatch):
    html = (
        b"<html><head><style>b{}</style>"
        b"<script>var x=1;</script></head>"
        b"<body><h1>Title</h1><p>Hello <b>world</b></p></body></html>"
    )
    seen = _fake_requests(monkeypatch, html, "text/html; charset=utf-8")
    text = fetch_url("https://example.org/page")
    assert "https://example.org/page" in seen
    assert "Title" in text
    assert "Hello" in text
    assert "world" in text
    assert "var x" not in text
    assert "<p>" not in text


def test_fetch_url_plain_text_passthrough_and_clip(monkeypatch):
    _fake_requests(monkeypatch, b"a" * 30, "text/plain")
    assert fetch_url("http://example.org/x") == "a" * 30
    assert "more characters" in fetch_url("http://example.org/x", max_chars=10)


def test_fetch_url_leaves_structured_text_alone(monkeypatch):
    _fake_requests(monkeypatch, b'{"a": 1, "b": [1, 2]}', "application/json")
    assert fetch_url("http://example.org/data.json") == '{"a": 1, "b": [1, 2]}'


def test_fetch_url_rejects_non_http():
    with pytest.raises(ValueError, match="http"):
        fetch_url("file:///etc/passwd")


# Models (llama3.2 among them) send null for optional arguments they do
# not set; the tools must take that as "the default", and Strands' argument
# validation must let it through, or the call fails before the tool runs.


def test_optional_arguments_accept_none(tmp_path, monkeypatch):
    (tmp_path / "a.txt").write_text("alpha\nbeta\n")
    assert read_file(str(tmp_path / "a.txt"), max_chars=None) == "alpha\nbeta\n"
    assert "a.txt" in list_files(str(tmp_path), pattern=None)
    monkeypatch.chdir(tmp_path)
    assert "a.txt" in list_files(None, None)  # None directory is ".", like the default
    assert "a.txt:2: beta" in search_files(
        "beta", str(tmp_path), file_glob=None, max_matches=None
    )
    assert current_time(None).count("-") >= 2
    assert regex_extract(r"\w+", "a b c", max_matches=None) == "a\nb\nc"
    _fake_requests(monkeypatch, b"hello", "text/plain")
    assert fetch_url("http://example.org/x", max_chars=None) == "hello"


def test_agent_tool_call_with_null_optional_argument_succeeds(scripted_model):
    scripted_model.call("current_time", timezone=None).say("it is now")
    seg = LLMAgent(source="scripted", model="test", tools=["current_time"])
    assert list(seg.transform(iter(["what time is it?"]))) == ["it is now"]
    [(name, status, _)] = scripted_model.tool_results
    assert (name, status) == ("current_time", "success")


# --- notepad ----------------------------------------------------------
#
# The pad is the agent's Strands state, so the tools are exercised with
# a stand-in context that carries one, and once through real agents to
# show where the state actually lives.


def _pad_context() -> SimpleNamespace:
    return SimpleNamespace(agent=SimpleNamespace(state=AgentState()))


def _notepad():
    return (
        common_tools.save_note,
        common_tools.read_notes,
        common_tools.summarize_notes,
        common_tools.update_note,
        common_tools.delete_note,
        common_tools.clear_notes,
    )


def test_notepad_pads_are_per_agent_state():
    save, read, *_ = _notepad()
    a, b = _pad_context(), _pad_context()
    assert read(tool_context=a) == "(no notes saved yet)"
    assert save("first", tool_context=a) == "(note 1 saved)"
    assert save("second", tool_context=a) == "(note 2 saved)"
    assert read(tool_context=a) == "1. first\n2. second"
    assert read(tool_context=b) == "(no notes saved yet)"
    # what is stored is plain JSON — a session manager can persist it
    assert json.loads(json.dumps(a.agent.state.get())) == a.agent.state.get()


def test_notepad_summarize_previews_without_full_notes():
    save, _read, summarize, *_ = _notepad()
    ctx = _pad_context()
    assert summarize(tool_context=ctx) == "(no notes saved yet)"
    save("short", tool_context=ctx)
    save("first line\nsecond line that should not appear", tool_context=ctx)
    save("x" * 200, tool_context=ctx)
    summary = summarize(tool_context=ctx)
    assert summary.startswith("3 notes, ")
    assert "1. (5 chars) short" in summary
    assert "2. (45 chars) first line…" in summary
    assert "second line" not in summary
    assert "3. (200 chars) " + "x" * 80 + "…" in summary


def test_notepad_update_and_delete_keep_numbers_stable():
    save, read, _summarize, update, delete, _clear = _notepad()
    ctx = _pad_context()
    save("first", tool_context=ctx)
    save("second", tool_context=ctx)
    save("third", tool_context=ctx)
    assert update(2, "revised", tool_context=ctx) == "(note 2 updated)"
    assert delete(1, tool_context=ctx) == "(note 1 deleted)"
    assert read(tool_context=ctx) == "2. revised\n3. third"
    assert save("fourth", tool_context=ctx) == "(note 4 saved)"
    with pytest.raises(ValueError, match=r"no note 1 \(saved notes: 2, 3, 4\)"):
        delete(1, tool_context=ctx)
    with pytest.raises(ValueError, match="no note 9"):
        update(9, "nope", tool_context=ctx)


def test_notepad_clear_expunges_everything():
    save, read, _summarize, _update, delete, clear = _notepad()
    ctx = _pad_context()
    assert clear(tool_context=ctx) == "(the pad was already empty)"
    save("a", tool_context=ctx)
    save("b", tool_context=ctx)
    assert clear(tool_context=ctx) == "(2 notes deleted)"
    assert read(tool_context=ctx) == "(no notes saved yet)"
    with pytest.raises(ValueError, match="saved notes: none"):
        delete(1, tool_context=ctx)
    assert save("c", tool_context=ctx) == "(note 3 saved)"  # numbering continues


def test_notepad_lives_in_the_conversation_for_chat(scripted_model):
    scripted_model.script.extend(
        [("tool", "save_note", {"note": "remember me"}), ("text", "noted")]
    )
    chat = AgentChat(source="scripted", model="test", tools=["save_note"])
    assert chat.ask("save this") == "noted"
    assert chat.agent.state.get("notepad")["notes"] == {"1": "remember me"}
    other = AgentChat(source="scripted", model="test", tools=["save_note"])
    assert other.agent.state.get("notepad") is None


def test_notepad_carries_across_items_in_a_segment(scripted_model):
    """A fresh agent per item, but one pad per segment: notes saved on
    one item are readable on the next."""
    scripted_model.script.extend(
        [
            ("tool", "save_note", {"note": "from item one"}),
            ("text", "saved"),
            ("tool", "read_notes", {}),
            ("text", "read"),
        ]
    )
    seg = LLMAgent(source="scripted", model="test", tools=["save_note", "read_notes"])
    assert list(seg.transform(iter(["a", "b"]))) == ["saved", "read"]
    assert scripted_model.calls == [1, 3, 1, 3]  # conversations did not carry over
    read_result = next(
        c["toolResult"]
        for m in scripted_model.prompts[-1]
        for c in m["content"]
        if "toolResult" in c
    )
    assert "1. from item one" in json.dumps(read_result["content"])


# --- end-to-end through the agent loop --------------------------------


def test_agent_calls_calculator_by_name(scripted_model):
    scripted_model.script.extend(
        [("tool", "calculator", {"expression": "6 * 7"}), ("text", "the answer is 42")]
    )
    chat = AgentChat(source="scripted", model="test", tools=["calculator"])
    assert chat.ask("what is 6 times 7?") == "the answer is 42"
    results = [
        c["toolResult"]
        for m in chat.messages
        for c in m["content"]
        if "toolResult" in c
    ]
    assert results
    assert results[0]["status"] == "success"
    assert "42" in json.dumps(results[0]["content"])
